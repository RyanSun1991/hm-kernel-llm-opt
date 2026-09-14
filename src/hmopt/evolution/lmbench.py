"""Convert explicit, independent lmbench suite pairs into a frozen A/B report.

The relay's digest and automatic ``prev_total_xlsx`` are discovery aids only.
Each statistical pair here is two independently collected suite files. Raw
``value0..N`` repeats within a suite are averaged, never promoted into extra
independent pairs. The manifest declares collection provenance; it does not
authenticate a device, a flashed image, or the independence of the runs.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
import statistics
import zipfile
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, ConfigDict, Field, field_validator

from .reports import register_report_conversion
from .store import digest
from .validation import ABReport, ExperimentArm, Identifier, Measurement, Sha256

_CONTEXT_FIELDS = (
    "repo_revision",
    "project_revisions",
    "image_sha256",
    "device_id",
    "workload_id",
    "workload_config_sha256",
    "environment_sha256",
)
_KEY_FIELDS = ("system", "tool", "metric", "command", "units")
_MAX_FILE_BYTES = 16 * 1024 * 1024
_MAX_TOTAL_BYTES = 64 * 1024 * 1024
_MAX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
_MAX_ROWS = 10000
_MAX_COLUMNS = 1024


class _StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)


class LmbenchMetric(_StrictModel):
    """Exact native row identity mapped to one metric in the approved policy."""

    name: Identifier
    system: Identifier
    tool: Identifier
    metric: Identifier
    command: Identifier
    units: Identifier


class LmbenchProfile(_StrictModel):
    """Hash this complete row mapping into the plan before plan review."""

    kind: Literal["lmbench_paired_suites"] = "lmbench_paired_suites"
    metrics: list[LmbenchMetric] = Field(min_length=1, max_length=20)


class LmbenchRun(_StrictModel):
    run_token: Identifier
    xlsx_path: Identifier
    xlsx_sha256: Sha256
    repo_revision: Identifier
    project_revisions: dict[Identifier, Identifier] = Field(default_factory=dict)
    image_sha256: Sha256
    device_id: Identifier
    workload_id: Identifier
    workload_config_sha256: Sha256
    environment_sha256: Sha256

    @field_validator("xlsx_sha256", "image_sha256", "workload_config_sha256", "environment_sha256")
    @classmethod
    def normalize_hash(cls, value: str) -> str:
        return value.lower()


class LmbenchPair(_StrictModel):
    pair_id: Identifier
    baseline: LmbenchRun
    candidate: LmbenchRun


class LmbenchManifest(_StrictModel):
    schema_version: Literal[1] = 1
    kind: Literal["lmbench_paired_suites"] = "lmbench_paired_suites"
    candidate_id: Identifier
    plan_digest: Sha256
    implementation_digest: Sha256
    baseline: ExperimentArm
    candidate: ExperimentArm
    metrics: list[LmbenchMetric] = Field(min_length=1, max_length=20)
    pairs: list[LmbenchPair] = Field(min_length=3, max_length=100)
    functional_passed: bool


def _error(message: str):
    from .service import GateError

    return GateError(message)


def _key(values) -> tuple[str, ...]:
    # Same whitespace normalization as the relay's _match_key, plus units.
    return tuple(re.sub(r"\s+", " ", str(value or "")).strip() for value in values)


def _artifact_path(value: str, artifacts_root: Path | None) -> Path:
    raw = Path(value)
    if ".." in raw.parts:
        raise _error("lmbench artifact paths must not contain parent traversal")
    if artifacts_root is None:
        if not raw.is_absolute():
            raise _error("lmbench xlsx_path must be absolute without an artifacts_root")
        result = raw.resolve(strict=True)
    else:
        root = Path(artifacts_root).resolve(strict=True)
        if not root.is_dir():
            raise _error("lmbench artifacts_root must be a directory")
        lexical = raw if raw.is_absolute() else root / raw
        try:
            lexical.relative_to(root)
        except ValueError as exc:
            raise _error("lmbench artifact is outside the configured artifacts_root") from exc
        result = lexical.resolve(strict=True)
        if not result.is_relative_to(root):
            raise _error("lmbench symlink target is outside the configured artifacts_root")
        # Refuse symlink/junction routing even when its destination stays inside.
        for component in (lexical, *lexical.parents):
            if component == root:
                break
            if component.is_symlink() or getattr(component, "is_junction", lambda: False)():
                raise _error("lmbench artifacts under artifacts_root must not use symlinks")
    if result.suffix.lower() != ".xlsx" or not result.is_file():
        raise _error("lmbench source must be a regular raw .xlsx file, not a digest")
    return result


def _read_workbook(content: bytes, metrics: list[LmbenchMetric]) -> dict:
    """Read the hashed byte snapshot, retaining selected raw values for audit.

    This deliberately does not call the relay's parse_total: that parser drops
    raw samples and tolerates missing values for its diagnostic digest. Column
    names and whitespace matching follow the same native workbook contract.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise _error(
            "lmbench conversion requires openpyxl; install it in the collector runtime"
        ) from exc

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 1000
                or sum(item.file_size for item in entries) > _MAX_UNCOMPRESSED_BYTES
            ):
                raise _error("lmbench workbook exceeds the expanded archive limit")
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        raise _error(f"Invalid lmbench workbook: {exc}") from exc
    try:
        if "result" not in workbook.sheetnames:
            raise _error("lmbench workbook must contain the native result sheet")
        sheet = workbook["result"]
        if (sheet.max_row or 0) > _MAX_ROWS + 1 or (sheet.max_column or 0) > _MAX_COLUMNS:
            raise _error("lmbench worksheet exceeds the row/column limit")
        rows = sheet.iter_rows()
        header_row = next(rows, None)
        if not header_row:
            raise _error("lmbench workbook has no header")
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in header_row]
        named = [name for name in headers if name]
        if len(named) != len(set(named)):
            raise _error("Duplicate lmbench column header (including duplicate raw samples)")
        index = {name: position for position, name in enumerate(headers) if name}
        if not set(_KEY_FIELDS).issubset(index):
            raise _error("lmbench workbook lacks complete native metric identity columns")
        samples = sorted(
            (name for name in index if re.fullmatch(r"value\d+", name)),
            key=lambda name: int(name[5:]),
        )
        if not samples or samples != [f"value{i}" for i in range(len(samples))]:
            raise _error(
                "lmbench requires contiguous raw value0..N columns; summaries are insufficient"
            )
        wanted = {
            _key(getattr(metric, field) for field in _KEY_FIELDS): metric for metric in metrics
        }
        found = {}
        for row_number, row in enumerate(rows, start=2):
            if row_number > _MAX_ROWS + 1 or len(row) > _MAX_COLUMNS:
                raise _error("lmbench worksheet exceeds the row/column limit")
            key = _key(row[index[field]].value for field in _KEY_FIELDS)
            if key not in wanted:
                continue
            metric = wanted[key]
            if metric.name in found:
                raise _error(f"Duplicate lmbench metric row: {metric.name}")
            raw_values = []
            for name in samples:
                cell = row[index[name]]
                value = cell.value
                if cell.data_type == "f" or type(value) not in (int, float):
                    raise _error(f"Missing or nonnumeric raw lmbench sample {metric.name}/{name}")
                try:
                    number = float(value)
                except (ValueError, OverflowError) as exc:
                    raise _error("Invalid raw lmbench sample") from exc
                if not math.isfinite(number) or number < 0:
                    raise _error("Raw lmbench samples must be finite and nonnegative")
                raw_values.append(number)
            mean = statistics.mean(raw_values)
            if not math.isfinite(mean):
                raise _error("lmbench suite mean is not finite")
            found[metric.name] = {
                "identity": dict(zip(_KEY_FIELDS, key)),
                "row": row_number,
                "sample_columns": samples,
                "raw_samples": raw_values,
                "suite_mean": mean,
            }
        missing = sorted({metric.name for metric in metrics} - found.keys())
        if missing:
            raise _error("Missing lmbench metric identity or units: " + ", ".join(missing))
        return found
    finally:
        workbook.close()


def convert_lmbench_report(
    service,
    candidate_id: str,
    manifest: LmbenchManifest,
    *,
    artifacts_root: Path | None = None,
) -> dict:
    """Read explicitly selected suite artifacts and register a canonical report.

    With artifacts_root set, all paths must resolve within that server-controlled
    directory without symlink routing. Otherwise operator paths must be absolute.
    Full selected raw samples and file hashes are preserved in SQLite evidence;
    the original workbooks should also be retained in the operator artifact store.
    """
    manifest = LmbenchManifest.model_validate(manifest.model_dump(mode="python"))
    data = service.store.read("candidate", candidate_id)["data"]
    if data["stage"] != "code_approved" or data.get("validation"):
        raise _error("lmbench conversion requires an unsealed code-approved candidate")
    if (
        manifest.candidate_id != candidate_id
        or manifest.plan_digest != data["plan_digest"]
        or manifest.implementation_digest != data["implementation_digest"]
    ):
        raise _error("lmbench manifest is not bound to the approved plan and implementation")
    policy = data["plan"]["validation"]
    if policy.get("kind") == "correctness":
        raise _error("lmbench measurements cannot replace correctness acceptance")
    profile = LmbenchProfile(metrics=manifest.metrics)
    if policy.get("measurement_profile_sha256") != digest(profile.model_dump(mode="json")):
        raise _error("lmbench metric profile was not frozen in the approved plan")
    if manifest.baseline.measurements or manifest.candidate.measurements:
        raise _error("lmbench manifest arms must not contain prefilled measurements")
    if (
        manifest.baseline.repo_revision != data["candidate"]["repo_revision"]
        or manifest.candidate.repo_revision != data["implementation"]["revision"]
    ):
        raise _error("lmbench manifest revisions do not match reviewed source")
    if manifest.baseline.repo_revision == manifest.candidate.repo_revision:
        raise _error("lmbench arms require distinct repository revisions")
    if manifest.baseline.image_sha256 == manifest.candidate.image_sha256:
        raise _error("lmbench arms require distinct image hashes")
    for arm in (manifest.baseline, manifest.candidate):
        for field in ("device_id", "workload_id", "workload_config_sha256", "environment_sha256"):
            if getattr(arm, field) != policy[field]:
                raise _error(f"lmbench manifest {field} differs from the frozen plan")
    rules = {rule["name"]: rule for rule in policy["metrics"]}
    if len({metric.name for metric in manifest.metrics}) != len(manifest.metrics):
        raise _error("Duplicate lmbench metric mapping names")
    identities = [
        _key(getattr(metric, field) for field in _KEY_FIELDS) for metric in manifest.metrics
    ]
    if len(set(identities)) != len(identities):
        raise _error("Duplicate lmbench native metric identity")
    if set(rules) != {metric.name for metric in manifest.metrics}:
        raise _error("lmbench metrics must match every frozen primary metric and guardrail")
    for metric in manifest.metrics:
        if metric.units != rules[metric.name]["unit"]:
            raise _error("lmbench metric units differ from the frozen plan")
    if len(manifest.pairs) < policy["minimum_pairs"]:
        raise _error("lmbench has fewer independent suite pairs than the approved minimum")

    seen_pairs, seen_tokens, seen_files, seen_hashes = set(), set(), set(), set()
    measurements = {"baseline": [], "candidate": []}
    artifacts = []
    total_bytes = 0
    for pair in manifest.pairs:
        if pair.pair_id in seen_pairs:
            raise _error("Duplicate lmbench suite pair identity")
        seen_pairs.add(pair.pair_id)
        for label in ("baseline", "candidate"):
            run = getattr(pair, label)
            arm = getattr(manifest, label)
            if any(getattr(run, field) != getattr(arm, field) for field in _CONTEXT_FIELDS):
                raise _error("lmbench run provenance differs from its frozen experiment arm")
            if run.run_token in seen_tokens:
                raise _error("Duplicate lmbench run_token; suite runs must be independent")
            seen_tokens.add(run.run_token)
            try:
                path = _artifact_path(run.xlsx_path, artifacts_root)
                if path in seen_files:
                    raise _error("Duplicate lmbench xlsx file reused across suite pairs")
                seen_files.add(path)
                with path.open("rb") as handle:
                    content = handle.read(_MAX_FILE_BYTES + 1)
            except OSError as exc:
                raise _error(f"Cannot read lmbench artifact: {exc}") from exc
            total_bytes += len(content)
            if len(content) > _MAX_FILE_BYTES or total_bytes > _MAX_TOTAL_BYTES:
                raise _error("lmbench artifact byte limit exceeded")
            sha = hashlib.sha256(content).hexdigest()
            if sha != run.xlsx_sha256:
                raise _error("lmbench xlsx SHA-256 differs from the collection manifest")
            if sha in seen_hashes:
                raise _error("Duplicate lmbench workbook content reused across suite runs")
            seen_hashes.add(sha)
            extracted = _read_workbook(content, manifest.metrics)
            measurements[label].append(
                Measurement(
                    pair_id=pair.pair_id,
                    metrics={name: values["suite_mean"] for name, values in extracted.items()},
                )
            )
            artifacts.append(
                {
                    "pair_id": pair.pair_id,
                    "arm": label,
                    "run_token": run.run_token,
                    "path": str(path),
                    "sha256": sha,
                    "bytes": len(content),
                    "metrics": extracted,
                }
            )
    report = ABReport(
        candidate_id=candidate_id,
        implementation_revision=data["implementation"]["revision"],
        baseline=manifest.baseline.model_copy(update={"measurements": measurements["baseline"]}),
        candidate=manifest.candidate.model_copy(update={"measurements": measurements["candidate"]}),
        functional_passed=manifest.functional_passed,
        metrics=policy["metrics"],
        minimum_pairs=policy["minimum_pairs"],
    )
    return register_report_conversion(
        service,
        candidate_id,
        report,
        candidate_snapshot=data,
        source={
            "kind": "lmbench_paired_suites",
            "schema_version": 1,
            "aggregation": "suite_mean_of_all_raw_columns",
            "measurement_profile_sha256": digest(profile.model_dump(mode="json")),
            "manifest": manifest.model_dump(mode="json"),
            "artifacts": artifacts,
        },
    )
